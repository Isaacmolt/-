#!/usr/bin/env python3
"""
SheetCraft AI — QA 驗收腳本（測試小真）
驗收對象：Product 04 (Debt Payoff Planner) 與 Product 05 (50/30/20 Budget Template)
"""

import os
import re
import sys
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font


# ── 輔助工具 ──────────────────────────────────────────────

PASS = "\u2705 PASS"
FAIL = "\u274c FAIL"
WARN = "\u26a0\ufe0f WARN"

results_all = {}  # { product_label: [(check_name, status, detail)] }


def log(product, check, passed, detail=""):
    status = PASS if passed else FAIL
    results_all.setdefault(product, []).append((check, status, detail))
    print(f"  [{status}] {check}" + (f" — {detail}" if detail else ""))


def log_warn(product, check, detail=""):
    results_all.setdefault(product, []).append((check, WARN, detail))
    print(f"  [{WARN}] {check}" + (f" — {detail}" if detail else ""))


def check_balanced_parens(formula: str) -> bool:
    depth = 0
    in_string = False
    str_char = None
    for ch in formula:
        if in_string:
            if ch == str_char:
                in_string = False
            continue
        if ch in ('"', "'"):
            in_string = True
            str_char = ch
            continue
        if ch == '(':
            depth += 1
        elif ch == ')':
            depth -= 1
            if depth < 0:
                return False
    return depth == 0


def extract_sheet_refs(formula: str):
    """Extract sheet names referenced in formulas like 'SheetName'!A1 or SheetName!A1"""
    refs = re.findall(r"'([^']+)'!", formula)
    refs += re.findall(r"(?<!')(\w[\w ]*?)!", formula)
    return refs


# ── 通用檢查 ──────────────────────────────────────────────

def run_common_checks(filepath, label):
    print(f"\n{'='*70}")
    print(f"  品質驗收: {label}")
    print(f"  檔案路徑: {filepath}")
    print(f"{'='*70}")

    # 1. 檔案可正常開啟
    try:
        wb = load_workbook(filepath, data_only=False)
        log(label, "01 檔案可正常開啟 (load_workbook)", True)
    except Exception as e:
        log(label, "01 檔案可正常開啟 (load_workbook)", False, str(e))
        return None

    sheets = wb.sheetnames

    # 2. 分頁數量 >= 3
    log(label, "02 分頁數量 >= 3", len(sheets) >= 3,
        f"共 {len(sheets)} 個分頁: {sheets}")

    # 3. 無空白分頁 — 前 10 行至少有一個有值的 cell
    empty_sheets = []
    for sn in sheets:
        ws = wb[sn]
        has_value = False
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
            if any(v is not None for v in row):
                has_value = True
                break
        if not has_value:
            empty_sheets.append(sn)
    log(label, "03 無空白分頁 (前10行有值)", len(empty_sheets) == 0,
        f"空白分頁: {empty_sheets}" if empty_sheets else "全部分頁皆有內容")

    # 4. 有 Instructions 分頁
    log(label, "04 Instructions 分頁存在", "Instructions" in sheets,
        f"找到分頁: {sheets}")

    # 5. 公式數量 >= 5
    formulas = []
    for sn in sheets:
        ws = wb[sn]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append((sn, cell.coordinate, cell.value))
    log(label, "05 公式數量 >= 5", len(formulas) >= 5,
        f"共 {len(formulas)} 個公式")

    # 6. 公式語法 — 括號平衡 + 引用分頁存在
    paren_errors = []
    ref_errors = []
    for sn, coord, fml in formulas:
        if not check_balanced_parens(fml):
            paren_errors.append(f"{sn}!{coord}: {fml}")
        for ref_sheet in extract_sheet_refs(fml):
            if ref_sheet not in sheets:
                ref_errors.append(f"{sn}!{coord} 引用不存在的分頁 '{ref_sheet}': {fml}")

    log(label, "06a 公式括號平衡", len(paren_errors) == 0,
        f"{len(paren_errors)} 個錯誤" + (f": {paren_errors[:5]}" if paren_errors else ""))
    log(label, "06b 公式引用分頁存在", len(ref_errors) == 0,
        f"{len(ref_errors)} 個錯誤" + (f": {ref_errors[:5]}" if ref_errors else ""))

    # 7. 有樣式 (header bold / 自訂字體 / 填充色)
    has_bold = False
    has_fill = False
    has_custom_font = False
    for sn in sheets:
        ws = wb[sn]
        for row in ws.iter_rows(min_row=1, max_row=3):
            for cell in row:
                if cell.font and cell.font.bold:
                    has_bold = True
                if cell.font and cell.font.name and cell.font.name != "Calibri":
                    has_custom_font = True
                if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb and cell.fill.fgColor.rgb not in (None, "00000000"):
                    has_fill = True
    style_ok = has_bold or has_fill or has_custom_font
    detail_parts = []
    if has_bold: detail_parts.append("bold")
    if has_custom_font: detail_parts.append("自訂字體")
    if has_fill: detail_parts.append("填充色")
    log(label, "07 有樣式 (bold/字體/填充)", style_ok,
        f"偵測到: {', '.join(detail_parts)}" if detail_parts else "未偵測到任何樣式")

    # 8. 有條件格式
    cond_fmt_sheets = []
    for sn in sheets:
        ws = wb[sn]
        if ws.conditional_formatting._cf_rules:
            cond_fmt_sheets.append(sn)
    log(label, "08 有條件格式", len(cond_fmt_sheets) > 0,
        f"分頁: {cond_fmt_sheets}" if cond_fmt_sheets else "未找到條件格式")

    # 9. 有圖表
    chart_sheets = []
    for sn in sheets:
        ws = wb[sn]
        if ws._charts:
            chart_sheets.append(f"{sn} ({len(ws._charts)} 個)")
    log(label, "09 有圖表", len(chart_sheets) > 0,
        f"分頁: {chart_sheets}" if chart_sheets else "未找到圖表")

    # 10. 檔案大小 > 5KB
    fsize = os.path.getsize(filepath)
    log(label, "10 檔案大小 > 5KB", fsize > 5120,
        f"{fsize:,} bytes ({fsize/1024:.1f} KB)")

    # 11. 無 #REF!
    ref_errors_cells = []
    for sn in sheets:
        ws = wb[sn]
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if isinstance(cell.value, str) and "#REF!" in cell.value:
                    ref_errors_cells.append(f"{sn}!{cell.coordinate}")
    log(label, "11 無 #REF! 錯誤", len(ref_errors_cells) == 0,
        f"發現 {len(ref_errors_cells)} 個: {ref_errors_cells[:10]}" if ref_errors_cells else "無 #REF!")

    # 12. 分頁命名 — 無預設名
    default_names = [sn for sn in sheets if re.match(r'^Sheet\d*$', sn, re.IGNORECASE)]
    log(label, "12 分頁命名無預設名", len(default_names) == 0,
        f"預設名分頁: {default_names}" if default_names else "所有分頁皆已命名")

    return wb


# ── P04 特有檢查 ──────────────────────────────────────────

def run_p04_checks(wb, label):
    print(f"\n  --- P04 特有檢查 ---")
    sheets = wb.sheetnames

    # Snowball Method 分頁
    log(label, "P04-A Snowball Method 分頁存在", "Snowball Method" in sheets)

    # Avalanche Method 分頁
    log(label, "P04-B Avalanche Method 分頁存在", "Avalanche Method" in sheets)

    # Comparison 分頁
    log(label, "P04-C Comparison 分頁存在", "Comparison" in sheets)

    # Debt Inventory 至少 3 筆範例數據
    if "Debt Inventory" in sheets:
        ws = wb["Debt Inventory"]
        # 找到 header row，然後計算資料行數
        data_rows = 0
        header_found = False
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
            if not header_found:
                # 嘗試偵測 header（通常第一行或第二行有欄位名稱）
                if row and any(isinstance(v, str) and v.strip() for v in row if v is not None):
                    header_found = True
                    continue
            else:
                if any(v is not None for v in row):
                    data_rows += 1
        log(label, "P04-D Debt Inventory >= 3 筆範例數據", data_rows >= 3,
            f"找到 {data_rows} 筆資料行")
    else:
        log(label, "P04-D Debt Inventory >= 3 筆範例數據", False, "Debt Inventory 分頁不存在")


# ── P05 特有檢查 ──────────────────────────────────────────

def run_p05_checks(wb, label):
    print(f"\n  --- P05 特有檢查 ---")
    sheets = wb.sheetnames

    # Needs, Wants, Savings 分頁
    for name in ["Needs", "Wants", "Savings"]:
        log(label, f"P05-A '{name}' 分頁存在", name in sheets)

    # Dashboard 有收入輸入欄位
    if "Dashboard" in sheets:
        ws = wb["Dashboard"]
        income_found = False
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if isinstance(cell.value, str) and re.search(r'income|收入', cell.value, re.IGNORECASE):
                    income_found = True
                    break
            if income_found:
                break
        log(label, "P05-B Dashboard 有收入輸入欄位", income_found,
            "找到 income 相關欄位" if income_found else "未找到 income/收入 欄位")
    else:
        log(label, "P05-B Dashboard 有收入輸入欄位", False, "Dashboard 分頁不存在")

    # Monthly Summary 有 12 個月份
    if "Monthly Summary" in sheets:
        ws = wb["Monthly Summary"]
        months = {"jan", "feb", "mar", "apr", "may", "jun",
                  "jul", "aug", "sep", "oct", "nov", "dec",
                  "january", "february", "march", "april", "may", "june",
                  "july", "august", "september", "october", "november", "december"}
        found_months = set()
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and v.strip().lower() in months:
                    found_months.add(v.strip().lower())
        # Normalize to canonical month names
        canonical = set()
        month_map = {
            "jan": 1, "january": 1, "feb": 2, "february": 2,
            "mar": 3, "march": 3, "apr": 4, "april": 4,
            "may": 5, "jun": 6, "june": 6,
            "jul": 7, "july": 7, "aug": 8, "august": 8,
            "sep": 9, "september": 9, "oct": 10, "october": 10,
            "nov": 11, "november": 11, "dec": 12, "december": 12,
        }
        for m in found_months:
            canonical.add(month_map.get(m, m))
        log(label, "P05-C Monthly Summary 有 12 個月份", len(canonical) >= 12,
            f"找到 {len(canonical)} 個不同月份: {sorted(canonical)}")
    else:
        log(label, "P05-C Monthly Summary 有 12 個月份", False, "Monthly Summary 分頁不存在")


# ── 主程式 ────────────────────────────────────────────────

def main():
    files = [
        (
            "/home/user/-/products/04-debt-payoff-planner/Debt_Payoff_Planner_SheetCraft.xlsx",
            "P04 Debt Payoff Planner",
            run_p04_checks,
        ),
        (
            "/home/user/-/products/05-503020-budget/503020_Budget_Template_SheetCraft.xlsx",
            "P05 50/30/20 Budget Template",
            run_p05_checks,
        ),
    ]

    for filepath, label, extra_checks in files:
        wb = run_common_checks(filepath, label)
        if wb:
            extra_checks(wb, label)

    # ── 總結報告 ──
    print(f"\n{'='*70}")
    print("  品質驗收總結報告 — 測試小真")
    print(f"{'='*70}")

    total_pass = 0
    total_fail = 0
    total_warn = 0

    for product, checks in results_all.items():
        p = sum(1 for _, s, _ in checks if s == PASS)
        f = sum(1 for _, s, _ in checks if s == FAIL)
        w = sum(1 for _, s, _ in checks if s == WARN)
        total_pass += p
        total_fail += f
        total_warn += w
        status = "ALL PASS" if f == 0 else f"{f} FAILED"
        print(f"\n  {product}: {status} ({p} pass / {f} fail / {w} warn)")
        if f > 0:
            for name, s, detail in checks:
                if s == FAIL:
                    print(f"    {FAIL} {name}: {detail}")

    print(f"\n  總計: {total_pass} pass / {total_fail} fail / {total_warn} warn")
    verdict = "PASS — 品質驗收通過" if total_fail == 0 else "FAIL — 有項目未通過，需修正"
    print(f"\n  最終判定: {verdict}")
    print(f"{'='*70}\n")

    return 0 if total_fail == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
