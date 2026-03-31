"""
SheetCraft AI — Product Quality Checker (Agent: 品管阿瑞 Aray)
Validates all generated .xlsx products meet quality standards.
"""

import os
import sys
import glob

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

PRODUCTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "products")

QUALITY_CHECKS = [
    "has_multiple_sheets",
    "has_instructions_sheet",
    "has_formulas",
    "has_styled_headers",
    "no_empty_sheets",
    "reasonable_file_size",
]


def check_product(xlsx_path):
    filename = os.path.basename(xlsx_path)
    product_dir = os.path.basename(os.path.dirname(xlsx_path))
    results = {"file": filename, "product": product_dir, "passed": [], "failed": [], "warnings": []}

    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    except Exception as e:
        results["failed"].append(f"Cannot open file: {e}")
        return results

    # Check 1: Multiple sheets
    if len(wb.sheetnames) >= 3:
        results["passed"].append(f"Multiple sheets ({len(wb.sheetnames)} sheets)")
    else:
        results["failed"].append(f"Too few sheets ({len(wb.sheetnames)}). Minimum 3 expected.")

    # Check 2: Instructions sheet
    has_instructions = any("instruct" in name.lower() for name in wb.sheetnames)
    if has_instructions:
        results["passed"].append("Has Instructions sheet")
    else:
        results["warnings"].append("No Instructions sheet found")

    # Check 3: Has formulas
    formula_count = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
    if formula_count >= 5:
        results["passed"].append(f"Has formulas ({formula_count} found)")
    elif formula_count > 0:
        results["warnings"].append(f"Few formulas ({formula_count}). Expected 5+.")
    else:
        results["failed"].append("No formulas found")

    # Check 4: Styled headers
    has_styled = False
    for ws in wb.worksheets:
        cell = ws.cell(row=1, column=1)
        if cell.font and (cell.font.bold or cell.font.size and cell.font.size > 11):
            has_styled = True
            break
        cell = ws.cell(row=1, column=2)
        if cell.font and (cell.font.bold or cell.font.size and cell.font.size > 11):
            has_styled = True
            break
    if has_styled:
        results["passed"].append("Has styled headers")
    else:
        results["warnings"].append("Headers may not be styled")

    # Check 5: No empty sheets
    empty_sheets = []
    for ws in wb.worksheets:
        has_data = False
        for row in ws.iter_rows(max_row=10, max_col=10):
            for cell in row:
                if cell.value is not None:
                    has_data = True
                    break
            if has_data:
                break
        if not has_data:
            empty_sheets.append(ws.title)
    if empty_sheets:
        results["failed"].append(f"Empty sheets: {', '.join(empty_sheets)}")
    else:
        results["passed"].append("No empty sheets")

    # Check 6: File size
    file_size = os.path.getsize(xlsx_path)
    if file_size < 1000:
        results["failed"].append(f"File too small ({file_size} bytes)")
    elif file_size > 10_000_000:
        results["warnings"].append(f"File very large ({file_size / 1_000_000:.1f} MB)")
    else:
        results["passed"].append(f"File size OK ({file_size / 1000:.0f} KB)")

    wb.close()
    return results


def main():
    print("=" * 60)
    print("  SheetCraft AI — Quality Check (品管阿瑞)")
    print("=" * 60)

    xlsx_files = glob.glob(os.path.join(PRODUCTS_DIR, "*", "*.xlsx"))

    if not xlsx_files:
        print("  No .xlsx files found in products/")
        return 1

    all_passed = True
    for xlsx_path in sorted(xlsx_files):
        results = check_product(xlsx_path)
        product = results["product"]
        passed = len(results["passed"])
        failed = len(results["failed"])
        warnings = len(results["warnings"])

        status = "PASS" if failed == 0 else "FAIL"
        if failed > 0:
            all_passed = False

        print(f"\n  [{status}] {product}")
        for p in results["passed"]:
            print(f"    ✓ {p}")
        for w in results["warnings"]:
            print(f"    ⚠ {w}")
        for f in results["failed"]:
            print(f"    ✗ {f}")

    print("\n" + "=" * 60)
    overall = "ALL PASSED" if all_passed else "ISSUES FOUND"
    print(f"  Overall: {overall} ({len(xlsx_files)} products checked)")
    print("=" * 60)

    return 0 if all_passed else 1


if __name__ == "__main__":
    sys.exit(main())
