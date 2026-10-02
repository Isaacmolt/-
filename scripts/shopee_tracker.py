#!/usr/bin/env python3
"""
蝦皮分潤 — 營運追蹤表產生器（Excel / 可匯入 Google Sheets）

產生 shopee-affiliate/tracker.xlsx，7 個分頁：
    使用說明｜選品庫｜貼文日誌｜每日數據｜月結對帳｜收入模型｜KPI 儀表板
公式全部內建（SUMIFS / COUNTIFS / MIN），把後台數字填進「每日數據」就會自動算到月結與儀表板。
「月結對帳」會自動標示：達 500 可提領、達 5 萬需轉公司戶。

用法：
    python scripts/shopee_tracker.py
"""

import csv
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE_DIR = Path(__file__).resolve().parent.parent
MODULE_DIR = BASE_DIR / "shopee-affiliate"
PRODUCTS_PATH = MODULE_DIR / "products.csv"
CONFIG_PATH = MODULE_DIR / "config.json"
CONFIG_EXAMPLE_PATH = MODULE_DIR / "config.example.json"
OUTPUT_PATH = MODULE_DIR / "tracker.xlsx"

# 品牌色（沿用 OPERATIONS.md）
GREEN = "2D5F2D"
CREAM = "F5F0E8"
GOLD = "D4A574"
ORANGE = "EE4D2D"   # 蝦皮橘，只用在警示
WHITE = "FFFFFF"

HEADER_FONT = Font(name="Microsoft JhengHei", bold=True, color=WHITE, size=11)
HEADER_FILL = PatternFill("solid", fgColor=GREEN)
INPUT_FILL = PatternFill("solid", fgColor="FFF9E6")
CALC_FILL = PatternFill("solid", fgColor=CREAM)
WARN_FILL = PatternFill("solid", fgColor="FDE2DC")
OK_FILL = PatternFill("solid", fgColor="E3F1E3")
THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BODY_FONT = Font(name="Microsoft JhengHei", size=10)
TITLE_FONT = Font(name="Microsoft JhengHei", bold=True, size=14, color=GREEN)

DATA_ROWS = 400  # 每日數據 / 貼文日誌預留列數


def load_config():
    path = CONFIG_PATH if CONFIG_PATH.exists() else CONFIG_EXAMPLE_PATH
    return json.load(open(path, encoding="utf-8"))


def load_products():
    with open(PRODUCTS_PATH, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 28


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def body(ws, row, col, value, fill=None, fmt=None, bold=False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name="Microsoft JhengHei", size=10, bold=bold)
    cell.border = BORDER
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


# ── 分頁 1：使用說明 ───────────────────────────────────────────────
def sheet_readme(wb, cfg):
    ws = wb.active
    ws.title = "使用說明"
    ws["A1"] = f"{cfg['brand']['account_name']} — 蝦皮分潤營運追蹤表"
    ws["A1"].font = TITLE_FONT
    lines = [
        "",
        "顏色規則：黃底 = 你要填的；米底 = 公式自動算，不要動。",
        "",
        "每天 5 分鐘：",
        "1. 發完文 → 到「貼文日誌」登記日期、平台、商品、sub_id。",
        "2. 打開 affiliate.shopee.tw 後台 → 把當天的點擊、訂單、分潤填進「每日數據」。",
        "3. 看「KPI 儀表板」，本月目標達成率一眼就知道。",
        "",
        "每月 1 號：",
        "1. 後台會產生上月「對帳明細」→ 把金額填進「月結對帳」的「後台對帳金額」。",
        "2. 累積達 NT$500 就申請提領；收到勞報單後 5 個工作日內回傳。",
        "3. 如果「營業稅警示」變紅（單月 ≥ NT$50,000），要開始準備公司戶，詳見 05-請款稅務與合規.md。",
        "",
        "匯入 Google Sheets：檔案 → 匯入 → 上傳這個 xlsx，公式會保留。",
        "",
        "重跑 python scripts/shopee_tracker.py 會覆蓋這個檔案，所以開始填數據後請另存一份，或改用 Google Sheets 版本。",
    ]
    for i, t in enumerate(lines, 2):
        ws.cell(row=i, column=1, value=t).font = BODY_FONT
    ws.column_dimensions["A"].width = 100


# ── 分頁 2：選品庫 ─────────────────────────────────────────────────
def sheet_products(wb, products):
    ws = wb.create_sheet("選品庫")
    headers = ["ID", "利基", "分類", "商品名稱", "售價", "預估分潤%", "預估每單分潤", "賣點", "痛點", "受眾", "蝦皮網址", "分潤短連結", "狀態", "備註"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(headers))
    for r, p in enumerate(products, 2):
        body(ws, r, 1, p["id"], INPUT_FILL)
        body(ws, r, 2, p["niche"], INPUT_FILL)
        body(ws, r, 3, p["category"], INPUT_FILL)
        body(ws, r, 4, p["name"], INPUT_FILL)
        body(ws, r, 5, float(p["price_twd"]), INPUT_FILL, "#,##0")
        body(ws, r, 6, float(p["est_commission_pct"]) / 100, INPUT_FILL, "0.0%")
        body(ws, r, 7, f"=MIN(E{r}*F{r},500)", CALC_FILL, "#,##0")
        body(ws, r, 8, p["selling_points"].replace("|", "、"), INPUT_FILL)
        body(ws, r, 9, p["pain_point"], INPUT_FILL)
        body(ws, r, 10, p["target_audience"], INPUT_FILL)
        body(ws, r, 11, p["shopee_url"], INPUT_FILL)
        body(ws, r, 12, p["affiliate_link"], INPUT_FILL)
        body(ws, r, 13, p["status"], INPUT_FILL)
        body(ws, r, 14, p["notes"], INPUT_FILL)
    dv = DataValidation(type="list", formula1='"idea,ready,live,paused"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"M2:M{DATA_ROWS}")
    set_widths(ws, [7, 8, 10, 28, 8, 10, 12, 36, 24, 14, 30, 30, 8, 30])
    ws.freeze_panes = "E2"


# ── 分頁 3：貼文日誌 ───────────────────────────────────────────────
def sheet_post_log(wb):
    ws = wb.create_sheet("貼文日誌")
    headers = ["日期", "平台", "貼文類型", "商品 ID", "商品名稱", "sub_id", "貼文網址", "觸及", "互動數", "備註"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(headers))
    for r in range(2, DATA_ROWS + 2):
        body(ws, r, 1, None, INPUT_FILL, "yyyy-mm-dd")
        for c in (2, 3, 4, 6, 7, 10):
            body(ws, r, c, None, INPUT_FILL)
        body(ws, r, 5, f'=IF(D{r}="","",IFERROR(VLOOKUP(D{r},選品庫!A:D,4,FALSE),"找不到"))', CALC_FILL)
        body(ws, r, 8, None, INPUT_FILL, "#,##0")
        body(ws, r, 9, None, INPUT_FILL, "#,##0")
    dv = DataValidation(type="list", formula1='"Threads,Instagram,蝦皮影音,YouTube,其他"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{DATA_ROWS + 1}")
    set_widths(ws, [12, 11, 12, 9, 26, 24, 30, 9, 9, 24])
    ws.freeze_panes = "A2"


# ── 分頁 4：每日數據 ───────────────────────────────────────────────
def sheet_daily(wb):
    ws = wb.create_sheet("每日數據")
    headers = ["日期", "平台", "點擊", "訂單數", "訂單金額", "分潤金（後台預估）", "備註", "月份（自動）"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(headers))
    for r in range(2, DATA_ROWS + 2):
        body(ws, r, 1, None, INPUT_FILL, "yyyy-mm-dd")
        body(ws, r, 2, None, INPUT_FILL)
        body(ws, r, 3, None, INPUT_FILL, "#,##0")
        body(ws, r, 4, None, INPUT_FILL, "#,##0")
        body(ws, r, 5, None, INPUT_FILL, "#,##0")
        body(ws, r, 6, None, INPUT_FILL, "#,##0")
        body(ws, r, 7, None, INPUT_FILL)
        body(ws, r, 8, f'=IF(A{r}="","",TEXT(A{r},"yyyy-mm"))', CALC_FILL)
    dv = DataValidation(type="list", formula1='"全部,Threads,Instagram,蝦皮影音,YouTube"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{DATA_ROWS + 1}")
    set_widths(ws, [12, 10, 8, 8, 10, 16, 24, 12])
    ws.freeze_panes = "A2"


# ── 分頁 5：月結對帳 ───────────────────────────────────────────────
def sheet_monthly(wb, start):
    ws = wb.create_sheet("月結對帳")
    headers = ["月份", "點擊（自動）", "訂單（自動）", "預估分潤（自動）", "後台對帳金額", "差異", "累積未提領", "可提領？", "營業稅警示", "勞報單已回傳", "入帳日", "備註"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(headers))
    y, m = start.year, start.month
    for i in range(18):
        r = i + 2
        month = f"{y}-{m:02d}"
        body(ws, r, 1, month, CALC_FILL)
        body(ws, r, 2, f'=SUMIFS(每日數據!C:C,每日數據!H:H,A{r})', CALC_FILL, "#,##0")
        body(ws, r, 3, f'=SUMIFS(每日數據!D:D,每日數據!H:H,A{r})', CALC_FILL, "#,##0")
        body(ws, r, 4, f'=SUMIFS(每日數據!F:F,每日數據!H:H,A{r})', CALC_FILL, "#,##0")
        body(ws, r, 5, None, INPUT_FILL, "#,##0")
        body(ws, r, 6, f'=IF(E{r}="","",E{r}-D{r})', CALC_FILL, "#,##0;[Red]-#,##0")
        prev = f"N(G{r-1})" if i > 0 else "0"
        body(ws, r, 7, f'=IF(E{r}="","",{prev}+E{r}-IF(K{r}="",0,{prev}+E{r}))', CALC_FILL, "#,##0")
        body(ws, r, 8, f'=IF(E{r}="","",IF(G{r}>=500,"✅ 可提領","未達 500，累積到下月"))', CALC_FILL)
        body(ws, r, 9, f'=IF(E{r}="","",IF(E{r}>=50000,"⚠️ 達 5 萬，需轉公司戶",IF(E{r}>=40000,"接近 5 萬，提前準備","OK")))', CALC_FILL)
        body(ws, r, 10, None, INPUT_FILL)
        body(ws, r, 11, None, INPUT_FILL, "yyyy-mm-dd")
        body(ws, r, 12, None, INPUT_FILL)
        m += 1
        if m > 12:
            m, y = 1, y + 1
    ws.conditional_formatting.add("I2:I19", FormulaRule(formula=['LEFT(I2,1)="⚠"'], fill=WARN_FILL))
    ws.conditional_formatting.add("I2:I19", FormulaRule(formula=['LEFT(I2,2)="接近"'], fill=PatternFill("solid", fgColor="FFF1CC")))
    ws.conditional_formatting.add("H2:H19", FormulaRule(formula=['LEFT(H2,1)="✅"'], fill=OK_FILL))
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("J2:J19")
    ws.cell(row=21, column=1, value="「累積未提領」邏輯：加上本月後台金額；當「入帳日」填入時視為已提領、歸零。「營業稅警示」依單月後台金額判斷。").font = BODY_FONT
    set_widths(ws, [10, 10, 10, 14, 14, 10, 12, 20, 24, 12, 12, 24])
    ws.freeze_panes = "B2"


# ── 分頁 6：收入模型 ───────────────────────────────────────────────
def sheet_model(wb, cfg):
    ws = wb.create_sheet("收入模型")
    ws["A1"] = "收入模型（改黃色格子，其他自動算）"
    ws["A1"].font = TITLE_FONT
    headers = ["假設", "保守", "基準", "積極", "說明"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=3, column=c, value=h)
    style_header(ws, 3, len(headers))
    inputs = [
        ("每月貼文數", 45, 60, 90, "Threads 2 篇/天 ≈ 60", "#,##0"),
        ("平均每篇觸及", 400, 800, 1800, "新帳號前 30 篇通常 100-500", "#,##0"),
        ("連結點擊率", 0.015, 0.025, 0.035, "留言區連結 1-4%", "0.0%"),
        ("7 天內轉換率", 0.02, 0.03, 0.04, "點了連結後下單的比例", "0.0%"),
        ("平均訂單金額", 450, 550, 650, "一張訂單整筆都算分潤", "#,##0"),
        ("平均分潤率", 0.015, 0.03, 0.04, "一般 1%、萬粉 5%、加碼 10%", "0.0%"),
        ("單筆分潤上限", 500, 500, 500, "蝦皮規定；YouTube 訂單 800", "#,##0"),
    ]
    for i, (label, a, b, c_, note, fmt) in enumerate(inputs):
        r = 4 + i
        body(ws, r, 1, label, bold=True)
        body(ws, r, 2, a, INPUT_FILL, fmt)
        body(ws, r, 3, b, INPUT_FILL, fmt)
        body(ws, r, 4, c_, INPUT_FILL, fmt)
        body(ws, r, 5, note)
    outputs = [
        ("月點擊", "={c}4*{c}5*{c}6", "#,##0"),
        ("月訂單", "={c}11*{c}7", "#,##0.0"),
        ("每單分潤", "=MIN({c}8*{c}9,{c}10)", "#,##0"),
        ("月分潤", "={c}12*{c}13", "#,##0"),
        ("投入時數（每篇 0.5h）", "={c}4*0.5", "#,##0"),
        ("時薪", "=IF({c}15=0,0,{c}14/{c}15)", "#,##0"),
    ]
    for i, (label, formula, fmt) in enumerate(outputs):
        r = 11 + i
        body(ws, r, 1, label, bold=True)
        for col, letter in ((2, "B"), (3, "C"), (4, "D")):
            body(ws, r, col, formula.format(c=letter), CALC_FILL, fmt, bold=(label == "月分潤"))
        body(ws, r, 5, "")
    goals = cfg.get("goals", {})
    body(ws, 18, 1, "目標（來自 config）", bold=True)
    for i, (k, v) in enumerate(goals.items()):
        r = 19 + i
        label = k.replace("month_", "第 ").replace("_commission_twd", " 個月")
        body(ws, r, 1, label)
        body(ws, r, 2, v, INPUT_FILL, "#,##0")
        body(ws, r, 3, f'=IF(B{r}=0,"",TEXT(C14/B{r},"0%")&" （基準情境達成率）")', CALC_FILL)
    set_widths(ws, [22, 12, 12, 12, 36])


# ── 分頁 7：KPI 儀表板 ─────────────────────────────────────────────
def sheet_dashboard(wb, cfg):
    ws = wb.create_sheet("KPI 儀表板", 0)
    ws["A1"] = f"{cfg['brand']['account_name']} — KPI 儀表板"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = '=TEXT(TODAY(),"yyyy-mm")'
    ws["A2"].font = Font(name="Microsoft JhengHei", size=10, color="888888")
    ws["B2"] = "← 本月（自動）"
    ws["B2"].font = Font(name="Microsoft JhengHei", size=10, color="888888")
    headers = ["指標", "本月", "上月", "累計", "說明"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=4, column=c, value=h)
    style_header(ws, 4, len(headers))
    this_m = '$A$2'
    last_m = 'TEXT(EDATE(TODAY(),-1),"yyyy-mm")'
    rows = [
        ("貼文數", f'=COUNTIFS(貼文日誌!A:A,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),貼文日誌!A:A,"<"&EDATE(DATE(YEAR(TODAY()),MONTH(TODAY()),1),1))',
                   f'=COUNTIFS(貼文日誌!A:A,">="&EDATE(DATE(YEAR(TODAY()),MONTH(TODAY()),1),-1),貼文日誌!A:A,"<"&DATE(YEAR(TODAY()),MONTH(TODAY()),1))',
                   '=COUNTA(貼文日誌!A2:A1000)', "創作者分級看「創作天數」，每天至少 1 篇"),
        ("點擊", f'=SUMIFS(每日數據!C:C,每日數據!H:H,{this_m})', f'=SUMIFS(每日數據!C:C,每日數據!H:H,{last_m})', '=SUM(每日數據!C2:C1000)', ""),
        ("訂單", f'=SUMIFS(每日數據!D:D,每日數據!H:H,{this_m})', f'=SUMIFS(每日數據!D:D,每日數據!H:H,{last_m})', '=SUM(每日數據!D2:D1000)', "等級 1 要日均 ≥1 單；等級 2 要日均 ≥10 單"),
        ("訂單金額", f'=SUMIFS(每日數據!E:E,每日數據!H:H,{this_m})', f'=SUMIFS(每日數據!E:E,每日數據!H:H,{last_m})', '=SUM(每日數據!E2:E1000)', ""),
        ("分潤金（預估）", f'=SUMIFS(每日數據!F:F,每日數據!H:H,{this_m})', f'=SUMIFS(每日數據!F:F,每日數據!H:H,{last_m})', '=SUM(每日數據!F2:F1000)', "以後台對帳為準"),
        ("轉換率（訂單/點擊）", '=IF(B6=0,"",B7/B6)', '=IF(C6=0,"",C7/C6)', '=IF(D6=0,"",D7/D6)', "健康值 2-5%"),
        ("每單平均分潤", '=IF(B7=0,"",B9/B7)', '=IF(C7=0,"",C9/C7)', '=IF(D7=0,"",D9/D7)', "太低 → 換高單價或加碼商品"),
        ("每篇貼文帶來分潤", '=IF(B5=0,"",B9/B5)', '=IF(C5=0,"",C9/C5)', '=IF(D5=0,"",D9/D5)', "衡量內容效率"),
    ]
    fmts = ["#,##0", "#,##0", "#,##0", "#,##0", "#,##0", "0.0%", "#,##0", "#,##0"]
    for i, (label, f1, f2, f3, note) in enumerate(rows):
        r = 5 + i
        body(ws, r, 1, label, bold=True)
        body(ws, r, 2, f1, CALC_FILL, fmts[i])
        body(ws, r, 3, f2, CALC_FILL, fmts[i])
        body(ws, r, 4, f3, CALC_FILL, fmts[i])
        body(ws, r, 5, note)
    # 目標達成率
    goals = cfg.get("goals", {})
    body(ws, 15, 1, "目標達成", bold=True)
    body(ws, 15, 2, "目標", bold=True)
    body(ws, 15, 3, "本月分潤", bold=True)
    body(ws, 15, 4, "達成率", bold=True)
    style_header(ws, 15, 4)
    for i, (k, v) in enumerate(goals.items()):
        r = 16 + i
        label = k.replace("month_", "第 ").replace("_commission_twd", " 個月目標")
        body(ws, r, 1, label)
        body(ws, r, 2, v, INPUT_FILL, "#,##0")
        body(ws, r, 3, "=B9", CALC_FILL, "#,##0")
        body(ws, r, 4, f"=IF(B{r}=0,0,C{r}/B{r})", CALC_FILL, "0%")
    ws.conditional_formatting.add("D16:D18", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=OK_FILL))
    body(ws, 20, 1, "提領狀態", bold=True)
    body(ws, 20, 2, '=IF(D9>=500,"✅ 累計已達 500，可到後台申請提領","未達 500，繼續累積")', CALC_FILL)
    ws.merge_cells("B20:E20")
    set_widths(ws, [22, 14, 14, 14, 40])


def main():
    cfg = load_config()
    products = load_products()
    wb = Workbook()
    sheet_readme(wb, cfg)
    sheet_products(wb, products)
    sheet_post_log(wb)
    sheet_daily(wb)
    sheet_monthly(wb, date.today().replace(day=1))
    sheet_model(wb, cfg)
    sheet_dashboard(wb, cfg)
    wb.save(OUTPUT_PATH)
    print("=" * 70)
    print("蝦皮分潤 — 營運追蹤表")
    print("=" * 70)
    print(f"已產生：{OUTPUT_PATH.relative_to(BASE_DIR)}")
    print(f"分頁：{', '.join(wb.sheetnames)}")
    print(f"選品庫已匯入 {len(products)} 個商品")


if __name__ == "__main__":
    main()
