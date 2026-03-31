#!/usr/bin/env python3
"""
品管阿瑞 — SheetCraft AI QA Script
Products 01, 02, 03 品質檢查
"""

import os
import re
from openpyxl import load_workbook
from openpyxl.chart import Reference


PRODUCTS = [
    {
        "name": "Product 01 — Monthly Budget Tracker",
        "path": "/home/user/-/products/01-monthly-budget-tracker/Monthly_Budget_Tracker_SheetCraft.xlsx",
    },
    {
        "name": "Product 02 — Annual Financial Overview",
        "path": "/home/user/-/products/02-annual-financial-overview/Annual_Financial_Overview_SheetCraft.xlsx",
    },
    {
        "name": "Product 03 — Freelancer Income Tracker",
        "path": "/home/user/-/products/03-freelancer-income-tracker/Freelancer_Income_Tracker_SheetCraft.xlsx",
    },
]

SEPARATOR = "=" * 72


def check_product(product):
    """Run all 12 QA checks on a single product file."""
    name = product["name"]
    path = product["path"]
    results = []  # list of (check_name, passed: bool, detail: str)
    wb = None

    # ── 1. File opens without error ──────────────────────────────────────
    try:
        wb = load_workbook(path, data_only=False)
        results.append(("1. 檔案可正常開啟", True, "load_workbook 成功"))
    except Exception as e:
        results.append(("1. 檔案可正常開啟", False, f"load_workbook 失敗: {e}"))
        # Cannot continue if file won't open
        return name, results

    sheet_names = wb.sheetnames

    # ── 2. At least 3 sheets ─────────────────────────────────────────────
    count = len(sheet_names)
    passed = count >= 3
    results.append(("2. 分頁數量 >= 3", passed, f"共 {count} 個分頁: {sheet_names}"))

    # ── 3. No blank sheets (first 10 rows must have at least 1 value) ───
    blank_sheets = []
    for sn in sheet_names:
        ws = wb[sn]
        has_value = False
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
            if any(c is not None for c in row):
                has_value = True
                break
        if not has_value:
            blank_sheets.append(sn)
    passed = len(blank_sheets) == 0
    detail = "全部分頁前 10 行皆有資料" if passed else f"空白分頁: {blank_sheets}"
    results.append(("3. 無空白分頁", passed, detail))

    # ── 4. Instructions sheet exists ─────────────────────────────────────
    has_instructions = any(s.lower().strip() == "instructions" for s in sheet_names)
    results.append(("4. 有 Instructions 分頁", has_instructions,
                     f"分頁列表: {sheet_names}"))

    # ── 5. At least 5 formulas ───────────────────────────────────────────
    formulas = []  # (sheet, cell, formula)
    for sn in sheet_names:
        ws = wb[sn]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append((sn, cell.coordinate, cell.value))
    formula_count = len(formulas)
    passed = formula_count >= 5
    results.append(("5. 公式數量 >= 5", passed, f"共 {formula_count} 個公式"))

    # ── 6. Formula syntax check ──────────────────────────────────────────
    syntax_errors = []
    for sn, coord, formula in formulas:
        # Check unmatched parentheses
        if formula.count("(") != formula.count(")"):
            syntax_errors.append((sn, coord, formula, "未閉合括號"))
        # Check sheet references — e.g. 'SheetName'!A1  or  SheetName!A1
        refs = re.findall(r"'([^']+)'!", formula) + re.findall(r"(?<!')(\w[\w ]*?)!", formula)
        for ref in refs:
            if ref not in sheet_names:
                syntax_errors.append((sn, coord, formula, f"引用不存在的分頁: {ref}"))
    passed = len(syntax_errors) == 0
    if passed:
        detail = "所有公式語法正確"
    else:
        detail = "語法錯誤:\n"
        for sn, coord, formula, err in syntax_errors:
            detail += f"      [{sn}] {coord}: {formula}  -> {err}\n"
    results.append(("6. 公式語法檢查", passed, detail))

    # ── 7. Header styling (bold or custom font) ─────────────────────────
    has_style = False
    style_detail = ""
    for sn in sheet_names:
        ws = wb[sn]
        for row in ws.iter_rows(min_row=1, max_row=3):
            for cell in row:
                if cell.font and (cell.font.bold or (cell.font.name and cell.font.name != "Calibri")
                                  or cell.font.size not in (None, 11)):
                    has_style = True
                    style_detail = f"[{sn}] {cell.coordinate}: bold={cell.font.bold}, font={cell.font.name}, size={cell.font.size}"
                    break
            if has_style:
                break
        if has_style:
            break
    results.append(("7. 有樣式/格式 (header)", has_style,
                     style_detail if has_style else "前 3 行無 bold / 自訂字體"))

    # ── 8. Conditional formatting ────────────────────────────────────────
    cf_sheets = []
    for sn in sheet_names:
        ws = wb[sn]
        if ws.conditional_formatting._cf_rules:
            cf_sheets.append(sn)
    passed = len(cf_sheets) > 0
    results.append(("8. 有條件格式", passed,
                     f"含條件格式的分頁: {cf_sheets}" if passed else "無任何條件格式"))

    # ── 9. Charts ────────────────────────────────────────────────────────
    chart_sheets = []
    for sn in sheet_names:
        ws = wb[sn]
        if ws._charts:
            chart_sheets.append((sn, len(ws._charts)))
    passed = len(chart_sheets) > 0
    if passed:
        detail = ", ".join(f"[{s}] {n} 個圖表" for s, n in chart_sheets)
    else:
        detail = "無任何圖表"
    results.append(("9. 有圖表", passed, detail))

    # ── 10. File size > 5 KB ─────────────────────────────────────────────
    size_bytes = os.path.getsize(path)
    size_kb = size_bytes / 1024
    passed = size_kb > 5
    results.append(("10. 檔案大小 > 5KB", passed, f"{size_kb:.1f} KB"))

    # ── 11. No #REF! errors ──────────────────────────────────────────────
    ref_errors = []
    for sn in sheet_names:
        ws = wb[sn]
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if val and isinstance(val, str) and "#REF!" in val:
                    ref_errors.append((sn, cell.coordinate, val))
    passed = len(ref_errors) == 0
    if passed:
        detail = "無 #REF! 錯誤"
    else:
        detail = "#REF! 錯誤:\n"
        for sn, coord, val in ref_errors:
            detail += f"      [{sn}] {coord}: {val}\n"
    results.append(("11. 無 #REF! 錯誤", passed, detail))

    # ── 12. No default sheet names like "Sheet1" ─────────────────────────
    default_names = [s for s in sheet_names if re.match(r"^Sheet\d+$", s)]
    passed = len(default_names) == 0
    results.append(("12. 分頁名稱無預設名稱", passed,
                     f"所有分頁: {sheet_names}" + (f" — 預設名稱: {default_names}" if default_names else "")))

    wb.close()
    return name, results


def main():
    print(SEPARATOR)
    print("  品管阿瑞 — SheetCraft AI 品質檢查報告")
    print("  Products 01, 02, 03")
    print(SEPARATOR)

    overall_all_pass = True

    for product in PRODUCTS:
        print(f"\n{SEPARATOR}")
        print(f"  {product['name']}")
        print(f"  {product['path']}")
        print(SEPARATOR)

        pname, results = check_product(product)
        all_pass = True
        for check_name, passed, detail in results:
            status = "PASS" if passed else "FAIL"
            icon = "[v]" if passed else "[X]"
            if not passed:
                all_pass = False
            print(f"  {icon} {status}  {check_name}")
            print(f"           {detail}")

        overall = "PASS" if all_pass else "FAIL"
        if not all_pass:
            overall_all_pass = False
        print(f"\n  >>> {pname} 最終判定: {overall}")

    print(f"\n{SEPARATOR}")
    final = "PASS" if overall_all_pass else "FAIL"
    print(f"  全部產品最終判定: {final}")
    print(SEPARATOR)


if __name__ == "__main__":
    main()
