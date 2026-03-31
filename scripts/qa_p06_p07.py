#!/usr/bin/env python3
"""
QA Script for SheetCraft AI — Products 06 & 07
審核阿德品管驗收腳本
"""

import os
import re
import sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

PASS = "\033[92mPASS\033[0m"
FAIL = "\033[91mFAIL\033[0m"
WARN = "\033[93mWARN\033[0m"

results_summary = {}


def log(product, check_id, ok, msg):
    status = PASS if ok else FAIL
    print(f"  [{status}] {check_id}: {msg}")
    results_summary.setdefault(product, []).append((check_id, ok, msg))


def check_brackets_balanced(formula):
    depth = 0
    for ch in formula:
        if ch == '(':
            depth += 1
        elif ch == ')':
            depth -= 1
        if depth < 0:
            return False
    return depth == 0


def extract_sheet_refs(formula):
    """Extract sheet name references like SheetName! from formulas."""
    # Match either 'Sheet Name'! or SimpleSheet! but not function( patterns
    # Pattern: word chars (optionally with spaces inside quotes) followed by !
    refs = []
    # Quoted sheet names: 'My Sheet'!
    refs.extend(re.findall(r"'([^']+)'!", formula))
    # Unquoted sheet names: must start with letter, only word chars
    # Use a lookbehind to avoid matching after ( which would be function names
    for m in re.finditer(r'(?<![A-Za-z(])([A-Za-z_]\w+)!', formula):
        refs.append(m.group(1))
    # Also catch sheet refs right after ( like COUNTA(Inventory!...)
    for m in re.finditer(r'\(([A-Za-z_]\w+)!', formula):
        refs.append(m.group(1))
    return refs


def collect_formulas(wb):
    formulas = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    formulas.append((ws.title, cell.coordinate, cell.value))
    return formulas


def run_common_checks(filepath, product_label):
    print(f"\n{'='*60}")
    print(f"  QA: {product_label}")
    print(f"  File: {filepath}")
    print(f"{'='*60}")

    # 1. File opens
    try:
        wb = load_workbook(filepath, data_only=False)
        log(product_label, "01-file-opens", True, "load_workbook succeeded")
    except Exception as e:
        log(product_label, "01-file-opens", False, f"load_workbook failed: {e}")
        return None

    sheet_names = wb.sheetnames
    print(f"  Sheets: {sheet_names}")

    # 2. At least 3 sheets
    ok = len(sheet_names) >= 3
    log(product_label, "02-sheet-count", ok,
        f"{len(sheet_names)} sheets (need >= 3)")

    # 3. No blank sheets (first 10 rows must have at least 1 non-empty cell)
    for sn in sheet_names:
        ws = wb[sn]
        has_data = False
        for row in ws.iter_rows(min_row=1, max_row=10):
            for cell in row:
                if cell.value is not None:
                    has_data = True
                    break
            if has_data:
                break
        log(product_label, f"03-not-blank[{sn}]", has_data,
            f"Sheet '{sn}' {'has' if has_data else 'MISSING'} data in first 10 rows")

    # 4. Instructions sheet
    has_instructions = any("instruct" in s.lower() for s in sheet_names)
    log(product_label, "04-instructions", has_instructions,
        f"Instructions sheet {'found' if has_instructions else 'NOT found'}")

    # 5. Formula count
    formulas = collect_formulas(wb)
    ok = len(formulas) >= 5
    log(product_label, "05-formula-count", ok,
        f"{len(formulas)} formulas (need >= 5)")

    # 6. Formula syntax
    bad_brackets = []
    bad_refs = []
    for sn, coord, f in formulas:
        if not check_brackets_balanced(f):
            bad_brackets.append(f"{sn}!{coord}: {f}")
        for ref_sheet in extract_sheet_refs(f):
            if ref_sheet not in sheet_names:
                bad_refs.append(f"{sn}!{coord} refs unknown sheet '{ref_sheet}': {f}")

    ok_br = len(bad_brackets) == 0
    log(product_label, "06a-brackets", ok_br,
        f"{'All balanced' if ok_br else 'UNBALANCED: ' + '; '.join(bad_brackets)}")
    ok_ref = len(bad_refs) == 0
    log(product_label, "06b-sheet-refs", ok_ref,
        f"{'All refs valid' if ok_ref else 'BAD REFS: ' + '; '.join(bad_refs)}")

    # 7. Header styling
    has_style = False
    for ws in wb.worksheets:
        for cell in ws[1]:  # first row
            if cell.font and (cell.font.bold or cell.font.size):
                has_style = True
                break
        if has_style:
            break
    log(product_label, "07-header-style", has_style,
        f"Header styling {'detected' if has_style else 'NOT detected'}")

    # 8. Conditional formatting
    has_cf = False
    cf_sheets = []
    for ws in wb.worksheets:
        if ws.conditional_formatting._cf_rules:
            has_cf = True
            cf_sheets.append(ws.title)
    log(product_label, "08-cond-format", has_cf,
        f"Conditional formatting in: {cf_sheets if has_cf else 'NONE'}")

    # 9. Charts
    has_chart = False
    chart_sheets = []
    for ws in wb.worksheets:
        if ws._charts:
            has_chart = True
            chart_sheets.append(ws.title)
    log(product_label, "09-charts", has_chart,
        f"Charts in: {chart_sheets if has_chart else 'NONE'}")

    # 10. File size
    fsize = os.path.getsize(filepath)
    ok = fsize > 5120
    log(product_label, "10-file-size", ok,
        f"{fsize:,} bytes ({'OK' if ok else 'TOO SMALL'}, need > 5KB)")

    # 11. No #REF!
    has_ref_err = False
    ref_err_locs = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if isinstance(val, str) and '#REF!' in val:
                    has_ref_err = True
                    ref_err_locs.append(f"{ws.title}!{cell.coordinate}")
    ok = not has_ref_err
    log(product_label, "11-no-ref-err", ok,
        f"{'No #REF! errors' if ok else '#REF! at: ' + ', '.join(ref_err_locs)}")

    # 12. No default sheet names
    default_names = {"Sheet", "Sheet1", "Sheet2", "Sheet3"}
    bad_names = [s for s in sheet_names if s in default_names]
    ok = len(bad_names) == 0
    log(product_label, "12-sheet-names", ok,
        f"{'All named properly' if ok else 'Default names found: ' + str(bad_names)}")

    return wb


def check_p06(wb):
    product = "P06"
    print(f"\n  --- P06 Extra Checks ---")
    sheet_names = wb.sheetnames

    # Find Monthly P&L sheet
    pl_sheet = None
    for sn in sheet_names:
        if "p&l" in sn.lower() or "pl" in sn.lower() or "profit" in sn.lower() or "loss" in sn.lower():
            pl_sheet = wb[sn]
            break
    if pl_sheet is None:
        # try broader match
        for sn in sheet_names:
            if "month" in sn.lower():
                pl_sheet = wb[sn]
                break

    if pl_sheet is None:
        log(product, "P06-pl-sheet", False, "No Monthly P&L sheet found")
        return

    log(product, "P06-pl-sheet", True, f"Found P&L sheet: '{pl_sheet.title}'")

    # Check sections: Revenue, COGS, Operating Expenses, Net Profit
    all_text = ""
    for row in pl_sheet.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                all_text += cell.value + "\n"

    sections = {
        "Revenue": False,
        "COGS": False,
        "Operating Expenses": False,
        "Net Profit": False,
    }
    for key in sections:
        for line in all_text.split("\n"):
            if key.lower() in line.lower():
                sections[key] = True
                break
    # Also check "Cost of Goods" as alternative for COGS
    if not sections["COGS"]:
        if "cost of goods" in all_text.lower():
            sections["COGS"] = True

    for key, found in sections.items():
        log(product, f"P06-section[{key}]", found,
            f"'{key}' {'found' if found else 'NOT found'} in P&L sheet")

    # Check 12 months columns
    month_names = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
                   "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    found_months = set()
    for row in pl_sheet.iter_rows(min_row=1, max_row=5):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                val = cell.value.strip()
                for m in month_names:
                    if m.lower() in val.lower():
                        found_months.add(m)

    ok = len(found_months) >= 12
    log(product, "P06-12-months", ok,
        f"Found {len(found_months)}/12 months: {sorted(found_months)}")

    # Check TOTAL column with SUM formulas
    total_col = None
    for row in pl_sheet.iter_rows(min_row=1, max_row=5):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and "total" in cell.value.lower():
                total_col = cell.column
                break
        if total_col:
            break

    if total_col is None:
        # Also check "Annual" or "Year"
        for row in pl_sheet.iter_rows(min_row=1, max_row=5):
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    if any(k in cell.value.lower() for k in ["annual", "year", "total", "fy"]):
                        total_col = cell.column
                        break
            if total_col:
                break

    if total_col is None:
        log(product, "P06-total-sum", False, "No TOTAL column found")
    else:
        sum_count = 0
        for row in pl_sheet.iter_rows(min_col=total_col, max_col=total_col):
            for cell in row:
                if isinstance(cell.value, str) and "SUM" in cell.value.upper():
                    sum_count += 1
        ok = sum_count > 0
        log(product, "P06-total-sum", ok,
            f"TOTAL column ({get_column_letter(total_col)}): {sum_count} SUM formulas found")


def check_p07(wb):
    product = "P07"
    print(f"\n  --- P07 Extra Checks ---")
    sheet_names = wb.sheetnames

    # Inventory sheet with >= 10 products
    inv_sheet = None
    for sn in sheet_names:
        if "inventor" in sn.lower() or "product" in sn.lower() or "stock" in sn.lower():
            # prefer "Inventory" specifically
            if "inventor" in sn.lower():
                inv_sheet = wb[sn]
                break
    if inv_sheet is None:
        for sn in sheet_names:
            if "inventor" in sn.lower() or "product" in sn.lower():
                inv_sheet = wb[sn]
                break

    if inv_sheet is None:
        log(product, "P07-inv-sheet", False, "No Inventory sheet found")
    else:
        # Count data rows (skip header)
        data_rows = 0
        for row in inv_sheet.iter_rows(min_row=2):
            has_data = any(c.value is not None for c in row)
            if has_data:
                data_rows += 1
        ok = data_rows >= 10
        log(product, "P07-inv-products", ok,
            f"Inventory sheet '{inv_sheet.title}': {data_rows} data rows (need >= 10)")

    # Stock Movement sheet with IN/OUT
    move_sheet = None
    for sn in sheet_names:
        if "movement" in sn.lower() or "transaction" in sn.lower() or "log" in sn.lower():
            move_sheet = wb[sn]
            break
    if move_sheet is None:
        for sn in sheet_names:
            if "stock" in sn.lower() and sn != (inv_sheet.title if inv_sheet else ""):
                move_sheet = wb[sn]
                break

    if move_sheet is None:
        log(product, "P07-movement", False, "No Stock Movement sheet found")
    else:
        all_vals = []
        for row in move_sheet.iter_rows():
            for cell in row:
                if cell.value:
                    all_vals.append(str(cell.value).upper())
        has_in = any("IN" == v or "IN" in v.split() for v in all_vals)
        has_out = any("OUT" == v or "OUT" in v.split() for v in all_vals)
        ok = has_in and has_out
        log(product, "P07-in-out", ok,
            f"Movement sheet '{move_sheet.title}': IN={'found' if has_in else 'MISSING'}, OUT={'found' if has_out else 'MISSING'}")

    # Suppliers sheet with >= 2 suppliers
    sup_sheet = None
    for sn in sheet_names:
        if "supplier" in sn.lower() or "vendor" in sn.lower():
            sup_sheet = wb[sn]
            break

    if sup_sheet is None:
        log(product, "P07-suppliers", False, "No Suppliers sheet found")
    else:
        data_rows = 0
        for row in sup_sheet.iter_rows(min_row=2):
            has_data = any(c.value is not None for c in row)
            if has_data:
                data_rows += 1
        ok = data_rows >= 2
        log(product, "P07-suppliers", ok,
            f"Suppliers sheet '{sup_sheet.title}': {data_rows} suppliers (need >= 2)")

    # SKU column exists somewhere
    sku_found = False
    sku_location = ""
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=3):
            for cell in row:
                if cell.value and isinstance(cell.value, str) and "sku" in cell.value.lower():
                    sku_found = True
                    sku_location = f"{ws.title}!{cell.coordinate}"
                    break
            if sku_found:
                break
        if sku_found:
            break
    log(product, "P07-sku", sku_found,
        f"SKU column {'found at ' + sku_location if sku_found else 'NOT found'}")


def main():
    p06_path = "/home/user/-/products/06-small-business-pl/Small_Business_PL_SheetCraft.xlsx"
    p07_path = "/home/user/-/products/07-inventory-manager/Inventory_Manager_SheetCraft.xlsx"

    print("=" * 60)
    print("  SheetCraft AI — QA Report (P06 & P07)")
    print("  審核阿德品管驗收報告")
    print("=" * 60)

    # P06
    wb06 = run_common_checks(p06_path, "P06")
    if wb06:
        check_p06(wb06)

    # P07
    wb07 = run_common_checks(p07_path, "P07")
    if wb07:
        check_p07(wb07)

    # Summary
    print(f"\n{'='*60}")
    print("  SUMMARY / 總結")
    print(f"{'='*60}")
    total_pass = 0
    total_fail = 0
    for product, checks in results_summary.items():
        p = sum(1 for _, ok, _ in checks if ok)
        f = sum(1 for _, ok, _ in checks if not ok)
        total_pass += p
        total_fail += f
        status = PASS if f == 0 else FAIL
        print(f"  {product}: {p} passed, {f} failed  [{status}]")

    print(f"\n  TOTAL: {total_pass} passed, {total_fail} failed")
    if total_fail == 0:
        print(f"\n  *** ALL CHECKS PASSED — 全部通過 ***")
    else:
        print(f"\n  *** {total_fail} ISSUE(S) FOUND — 發現 {total_fail} 個問題 ***")

    return 0 if total_fail == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
